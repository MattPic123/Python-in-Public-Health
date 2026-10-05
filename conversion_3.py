"""Converting the CSV to a formatted Excel workbook."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from dataset_utils import columns, source_file, load_records

OUTPUT_FILE = Path(__file__).with_name("maternal_health_risk.xlsx")


def main():
    if len(columns) != len(set(columns)):
        raise ValueError("Duplicate column names are not allowed.")

    records = load_records(source_file)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Risk Data"
    sheet.append(columns)

    for record in records:
        sheet.append([record[column] for column in columns])

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F1F1F")

    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, 4).number_format = "0.0#"
        sheet.cell(row_number, 5).number_format = "0.0#"

    column_widths = {
        "A": 10,
        "B": 15,
        "C": 15,
        "D": 10,
        "E": 12,
        "F": 12,
        "G": 14,
    }
    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 28

    table = Table(displayName="MaternalHealthRiskData", ref=f"A1:G{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A2"

    workbook.save(OUTPUT_FILE)

    # Confirm the saved file has the expected headers, unique columns, and rows.
    saved_workbook = load_workbook(OUTPUT_FILE, read_only=True, data_only=True)
    saved_sheet = saved_workbook["Risk Data"]
    saved_header = tuple(
        cell.value for cell in next(saved_sheet.iter_rows(min_row=1, max_row=1))
    )
    saved_row_count = sum(
        1 for _ in saved_sheet.iter_rows(min_row=2, values_only=True)
    )
    saved_workbook.close()

    if saved_header != tuple(columns) or len(set(saved_header)) != len(saved_header):
        raise ValueError("Saved workbook headers are incorrect or duplicated.")
    if saved_row_count != len(records):
        raise ValueError(
            f"Saved workbook has {saved_row_count} records; expected {len(records)}."
        )

    print(
        f"Created {OUTPUT_FILE.name}: "
        f"{saved_row_count} records, {len(saved_header)} unique columns."
    )


if __name__ == "__main__":
    main()